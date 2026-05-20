"""Audit export: live ads that become day-10 between May 15-18 with performance_score < 50.

Source rows = `ads` WHERE start_date BETWEEN 2026-05-05 AND 2026-05-08 AND performance_score < 50
(pulled once via mcp__supabase__execute_sql; result persisted to the tool-results file below).
Built because the v2 winner-alerts workflow dropped the score>=50 gate — this is the audit
of which sub-50-score ads now fall inside the processing window.
"""
import json
import re
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

SRC = "/Users/igordviniatin/.claude/projects/-Users-igordviniatin-Documents-thermoslim-platform/50b374d6-8061-4ee5-98a3-9e396b339d37/tool-results/mcp-supabase-execute_sql-1779115036605.txt"
OUT = "/Users/igordviniatin/Documents/thermoslim-platform/exports/low_score_under50_day10_2026-05-18.xlsx"

inner = json.loads(open(SRC).read())["result"]
rows = json.loads(re.search(r"\[.*\]", inner, re.S).group(0))

# 30 GetHookd-sourced ad fields, in the SELECT order.
COLS = [
    "id", "external_id", "brand_id", "brand_name", "brand_external_id", "platform",
    "display_format", "media_type", "title", "body", "cta_type", "cta_text",
    "landing_page", "link_description", "start_date", "end_date", "days_active",
    "active_in_library", "performance_score", "performance_score_title",
    "ad_spend_range_score", "ad_spend_range_score_title", "gender_audience",
    "age_audience_min", "age_audience_max", "share_url", "media_url",
    "thumbnail_url", "first_seen_at", "updated_at",
]
WIDTHS = {
    "title": 48, "body": 70, "landing_page": 50, "link_description": 40,
    "share_url": 55, "media_url": 55, "thumbnail_url": 55, "external_id": 20,
    "brand_external_id": 20, "performance_score_title": 18, "first_seen_at": 22,
    "updated_at": 22,
}

header_fill = PatternFill("solid", fgColor="1F4E78")
header_font = Font(bold=True, color="FFFFFF")


def style_header(ws):
    for c in ws[1]:
        c.fill = header_fill
        c.font = header_font
    ws.freeze_panes = "A2"


wb = Workbook()
wb.remove(wb.active)

# --- Summary sheet ---
summary = wb.create_sheet("Summary")
summary.append(["Metric", "Value"])
style_header(summary)
summary.append(["Total ads (score < 50, start 2026-05-05..05-08)", len(rows)])
summary.append(["", ""])
summary.append(["By start_date", "count"])
by_date = {}
for r in rows:
    by_date[r["start_date"]] = by_date.get(r["start_date"], 0) + 1
for d in sorted(by_date):
    summary.append([d, by_date[d]])
summary.append(["", ""])
summary.append(["By brand", "count"])
by_brand = {}
for r in rows:
    by_brand[r["brand_name"]] = by_brand.get(r["brand_name"], 0) + 1
for b in sorted(by_brand, key=lambda k: -by_brand[k]):
    summary.append([b, by_brand[b]])
summary.column_dimensions["A"].width = 46
summary.column_dimensions["B"].width = 12

# --- Ads sheet ---
ws = wb.create_sheet("Ads")
ws.append(COLS)
style_header(ws)
for r in rows:
    ws.append([r.get(c) for c in COLS])
for i, col in enumerate(COLS, 1):
    ws.column_dimensions[chr(64 + i) if i <= 26 else "A" + chr(64 + i - 26)].width = WIDTHS.get(col, 14)
last_col = ws.cell(row=1, column=len(COLS)).column_letter
ws.auto_filter.ref = f"A1:{last_col}{len(rows) + 1}"

os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print(f"Wrote {OUT}")
print(f"{len(rows)} rows, {len(COLS)} columns")
