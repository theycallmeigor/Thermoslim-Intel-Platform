"""Export: May-9 day-10 ads that SHOULD have fired a Discord alert but did not.

Winner set = `ads` WHERE start_date='2026-05-09' AND active_in_library=1
AND performance_score >= 70 (the DB/explore-endpoint score — the authoritative one).
33 rows. Run 3814 posted 0 of them because the workflow gate reads GetHookd's
`/ads/{id}` detail endpoint, which returns score 1/null for these same ads.

Extra column `api_score_exec3814` = the score `/ads/{id}` returned in execution 3814
(the value that routed each ad to Skip instead of Post). Lets the discrepancy be
verified row-by-row.
"""
import json
import re
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

SQL_SRC = "/Users/igordviniatin/.claude/projects/-Users-igordviniatin-Documents-thermoslim-platform/50b374d6-8061-4ee5-98a3-9e396b339d37/tool-results/mcp-supabase-execute_sql-1779201629532.txt"
EXEC_SRC = "/Users/igordviniatin/.claude/projects/-Users-igordviniatin-Documents-thermoslim-platform/50b374d6-8061-4ee5-98a3-9e396b339d37/tool-results/mcp-n8n-thermoslim-get_execution-1779201235495.txt"
OUT = "/Users/igordviniatin/Documents/thermoslim-platform/exports/missed_winners_may9_2026-05-19.xlsx"

# --- DB winner rows ---
inner = json.loads(open(SQL_SRC).read())["result"]
rows = json.loads(re.search(r"\[.*\]", inner, re.S).group(0))

# --- Map ad_db_id -> score that /ads/{id} returned in execution 3814 ---
exec_data = json.loads(open(EXEC_SRC).read())
api_score = {}
for run in exec_data["data"]["resultData"]["runData"].get("Build Result", []):
    for batch in run.get("data", {}).get("main", []):
        for item in (batch or []):
            j = item.get("json", {})
            api_score[j.get("ad_db_id")] = j.get("score")

# 30 GetHookd-sourced ad fields + the exec-3814 API score for cross-check.
COLS = [
    "id", "external_id", "brand_id", "brand_name", "brand_external_id", "platform",
    "display_format", "media_type", "title", "body", "cta_type", "cta_text",
    "landing_page", "link_description", "start_date", "end_date", "days_active",
    "active_in_library", "performance_score", "performance_score_title",
    "api_score_exec3814",
    "ad_spend_range_score", "ad_spend_range_score_title", "gender_audience",
    "age_audience_min", "age_audience_max", "share_url", "media_url",
    "thumbnail_url", "first_seen_at", "updated_at",
]
WIDTHS = {
    "title": 48, "body": 70, "landing_page": 50, "link_description": 40,
    "share_url": 55, "media_url": 55, "thumbnail_url": 55, "external_id": 20,
    "brand_external_id": 20, "performance_score_title": 18, "first_seen_at": 22,
    "updated_at": 22, "api_score_exec3814": 18,
}

header_fill = PatternFill("solid", fgColor="1F4E78")
header_font = Font(bold=True, color="FFFFFF")


def style_header(ws):
    for c in ws[1]:
        c.fill = header_fill
        c.font = header_font
    ws.freeze_panes = "A2"


wb = Workbook()
wb.remove(wb.active)

# --- Summary sheet ---
summary = wb.create_sheet("Summary")
summary.append(["Metric", "Value"])
style_header(summary)
summary.append(["Missed winners (DB score >= 70, start 2026-05-09)", len(rows)])
summary.append(["Suppressed because /ads/{id} returned null in exec 3814",
                 sum(1 for r in rows if api_score.get(r["id"]) is None)])
summary.append(["", ""])
summary.append(["By brand", "count"])
by_brand = {}
for r in rows:
    by_brand[r["brand_name"]] = by_brand.get(r["brand_name"], 0) + 1
for b in sorted(by_brand, key=lambda k: -by_brand[k]):
    summary.append([b, by_brand[b]])
summary.append(["", ""])
summary.append(["By DB performance_score", "count"])
by_score = {}
for r in rows:
    by_score[r["performance_score"]] = by_score.get(r["performance_score"], 0) + 1
for s in sorted(by_score, reverse=True):
    summary.append([s, by_score[s]])
summary.column_dimensions["A"].width = 52
summary.column_dimensions["B"].width = 12

# --- Ads sheet ---
ws = wb.create_sheet("Missed Winners")
ws.append(COLS)
style_header(ws)
for r in rows:
    r = dict(r)
    r["api_score_exec3814"] = api_score.get(r["id"])
    ws.append([r.get(c) for c in COLS])
for i, col in enumerate(COLS, 1):
    ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = WIDTHS.get(col, 14)
last_col = ws.cell(row=1, column=len(COLS)).column_letter
ws.auto_filter.ref = f"A1:{last_col}{len(rows) + 1}"

os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print(f"Wrote {OUT}")
print(f"{len(rows)} rows, {len(COLS)} columns")
