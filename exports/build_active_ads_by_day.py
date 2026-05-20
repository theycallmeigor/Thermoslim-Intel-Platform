import json
import re
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

SRC = "/Users/igordviniatin/.claude/projects/-Users-igordviniatin-Documents-thermoslim-platform/41bcf2c9-7e3c-4041-a5cb-7c653e5ce85b/tool-results/mcp-supabase-execute_sql-1779103216038.txt"
OUT = "/Users/igordviniatin/Documents/thermoslim-platform/exports/active_ads_by_day_2026-05-18.xlsx"

raw = open(SRC).read()
inner = json.loads(raw)["result"]
m = re.search(r"\[.*\]", inner, re.S)
rows = json.loads(m.group(0))

by_day = {}
for r in rows:
    by_day.setdefault(r["days_active"], []).append(r)

wb = Workbook()
wb.remove(wb.active)

cols = ["id", "external_id", "brand_name", "title", "start_date",
        "performance_score", "performance_score_title", "display_format", "share_url"]
header_fill = PatternFill("solid", fgColor="1F4E78")
header_font = Font(bold=True, color="FFFFFF")

# Summary sheet first
summary = wb.create_sheet("Summary")
summary.append(["days_active", "active_ads", "score>=50"])
for c in summary[1]:
    c.fill = header_fill
    c.font = header_font
for day in sorted(by_day):
    ads = by_day[day]
    s50 = sum(1 for a in ads if (a.get("performance_score") or 0) >= 50)
    summary.append([day, len(ads), s50])
summary.column_dimensions["A"].width = 14
summary.column_dimensions["B"].width = 14
summary.column_dimensions["C"].width = 12
summary.freeze_panes = "A2"

for day in sorted(by_day):
    ads = by_day[day]
    ws = wb.create_sheet(f"Day {day}")
    ws.append(cols)
    for c in ws[1]:
        c.fill = header_fill
        c.font = header_font
    for a in ads:
        ws.append([a.get(c) for c in cols])
    widths = [12, 20, 18, 48, 12, 10, 14, 14, 60]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{len(ads) + 1}"

os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print(f"Wrote {OUT}")
print(f"Total active ads: {len(rows)} across {len(by_day)} day-buckets (day {min(by_day)}-{max(by_day)})")
